"""DOCX renderer (python-docx). Uses real Word headings, caption styles and TOC fields."""
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING, WD_TAB_ALIGNMENT, \
    WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage

from model import runs

FONT = "Times New Roman"
MONO = "Consolas"
TEXT_W = 8.27 - 1.25 - 1.0   # inches


# ------------------------------------------------------------------ helpers

# Element orders from the WordprocessingML schema: insert children before their successors so
# Word never reports "unreadable content".
PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                 "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                 "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                 "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
PPR_AFTER_PBDR = ("w:shd",) + PPR_AFTER_SHD
TCPR_AFTER_BORDERS = ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                      "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
TCPR_AFTER_SHD = TCPR_AFTER_BORDERS[1:]
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                       "w:tblDescription", "w:tblPrChange")
SECT_AFTER_PGNUM = ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
                    "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
SETTINGS_AFTER_UPDATE = ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
                         "w:rsids", "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
                         "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
                         "w:captions", "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary",
                         "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator")


def insert_ordered(parent, child, successors):
    for tag in successors:
        if ":" in tag and not tag.startswith("w:"):
            continue
        found = parent.find(qn(tag))
        if found is not None:
            found.addprevious(child)
            return child
    parent.append(child)
    return child

def set_font(run_or_style, name=FONT, size=None, bold=None, italic=None, color=None):
    f = run_or_style.font
    f.name = name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style, "element") and \
        hasattr(run_or_style.element, "get_or_add_rPr") else run_or_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)


def add_runs(par, text, size=None, base_bold=False):
    for seg, style in runs(text):
        r = par.add_run(seg)
        if style == "code":
            set_font(r, MONO, (size or 12) - 1.5)
        else:
            set_font(r, FONT, size)
        if style == "bold" or base_bold:
            r.bold = True
        if style == "italic":
            r.italic = True
        if style == "fill":
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def para(doc_or_cell, text="", style=None, align=None, size=None, bold=False, space_after=None,
         space_before=None, keep_next=False, line=None):
    p = doc_or_cell.add_paragraph(style=style)
    if text:
        add_runs(p, text, size, bold)
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if keep_next:
        pf.keep_with_next = True
    if line is not None:
        pf.line_spacing = line
    return p


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    succ = TCPR_AFTER_SHD if el_pr.tag == qn("w:tcPr") else PPR_AFTER_SHD
    insert_ordered(el_pr, shd, succ)


def cell_borders(cell, val="single", sz=6, color="6B7280"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        b.append(e)
    insert_ordered(tcPr, b, TCPR_AFTER_BORDERS)


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    insert_ordered(tblPr, b, TBLPR_AFTER_BORDERS)


def field(par, instr, placeholder=""):
    r = par.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    r4 = par.add_run(placeholder)
    r5 = par.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)
    return r4


def fld_char(par, kind):
    r = par.add_run()
    c = OxmlElement("w:fldChar")
    c.set(qn("w:fldCharType"), kind)
    r._r.append(c)


def instr(par, text):
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = text
    r._r.append(it)


def pgnum(section, fmt, start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        insert_ordered(sectPr, pg, SECT_AFTER_PGNUM)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def footer_page(section):
    f = section.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = field(p, "PAGE", "1")
    for run in p.runs:
        set_font(run, FONT, 11)


def header_text(section, left, right):
    h = section.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0]
    p.text = ""
    p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(f"{left}\t{right}")
    set_font(r, FONT, 9.5, italic=True, color="6B7280")
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single")
    bt.set(qn("w:sz"), "4")
    bt.set(qn("w:color"), "9CA3AF")
    bt.set(qn("w:space"), "1")
    bdr.append(bt)
    insert_ordered(pPr, bdr, PPR_AFTER_PBDR)


def ensure_style(doc, name, base="Normal", ptype=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, ptype)
        st.base_style = doc.styles[base]
        return st


# ------------------------------------------------------------------ styles
def setup_styles(doc):
    normal = doc.styles["Normal"]
    set_font(normal, FONT, 12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def heading(name, size, align, before, after, outline):
        st = doc.styles[name]
        set_font(st, FONT, size, bold=True, italic=False, color="000000")
        st.paragraph_format.alignment = align
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15
        return st

    heading("Heading 1", 18, WD_ALIGN_PARAGRAPH.CENTER, 0, 24, 0)
    heading("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 12, 6, 1)
    heading("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 8, 4, 2)

    for name in ("Figure Caption", "Table Caption"):
        st = ensure_style(doc, name)
        set_font(st, FONT, 11)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.space_before = Pt(4)
        st.paragraph_format.space_after = Pt(12 if name == "Figure Caption" else 6)
        if name == "Table Caption":
            st.paragraph_format.keep_with_next = True

    st = ensure_style(doc, "Front Heading")
    set_font(st, FONT, 18, bold=True)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(20)

    for i, ind in ((1, 0), (2, 0.3), (3, 0.6)):
        st = ensure_style(doc, f"TOC {i}")
        set_font(st, FONT, 12, bold=(i == 1))
        st.paragraph_format.left_indent = Inches(ind)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.line_spacing = 1.2
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    st = ensure_style(doc, "Table of Figures")
    set_font(st, FONT, 11.5)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.2
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.left_indent = Inches(0.9)
    st.paragraph_format.first_line_indent = Inches(-0.9)
    st.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


# ------------------------------------------------------------------ blocks
def heading1(doc, label, title, page_break=True):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = page_break
    if label:
        r = p.add_run(label)
        set_font(r, FONT, 16, bold=True)
        r.add_break(WD_BREAK.LINE)
    r = p.add_run(title)
    set_font(r, FONT, 18, bold=True)
    return p


def contents_list(doc, entries, instr_text, style_for):
    """Write a Word TOC field, pre-filled with entries (Word refreshes it on open)."""
    if not entries:
        entries = [("toc", "Right-click and choose Update Field", 0, "", "")]
    first = True
    for i, (_, text, level, page, _k) in enumerate(entries):
        p = doc.add_paragraph(style=style_for(level))
        if first:
            fld_char(p, "begin")
            instr(p, instr_text)
            fld_char(p, "separate")
            first = False
        r = p.add_run(f"{text}\t{page}")
        set_font(r, FONT, None)
        if i == len(entries) - 1:
            fld_char(p, "end")


def caption_par(doc, label, num, text, style):
    p = doc.add_paragraph(style=style)
    if num:
        r = p.add_run(f"{label} {num}: ")
        set_font(r, FONT, 11, bold=True)
    add_runs(p, text, 11)
    return p


def picture_par(doc, path, width_in, max_h_in=7.6):
    with PILImage.open(path) as im:
        w, h = im.size
    width_in = min(width_in, TEXT_W)
    if width_in * h / w > max_h_in:
        width_in = max_h_in * w / h
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(path, width=Inches(width_in))
    return p


def placeholder_box(doc, height, title, hint, width=TEXT_W * 0.92):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    c.width = Inches(width)
    row = t.rows[0]
    row.height = Inches(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_borders(c, "dashed", 8, "9CA3AF")
    tcPr = c._tc.get_or_add_tcPr()
    shade(tcPr, "FAFAFA")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_font(r, FONT, 13, bold=True, color="4B5563")
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run(hint)
    set_font(r, FONT, 10.5, italic=True, color="4B5563")
    # keep the table row on one page with the caption
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)
    for pp in c.paragraphs:
        pp.paragraph_format.keep_with_next = True
    return t


def data_table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        scale = TEXT_W / sum(widths)
        widths = [w * scale for w in widths]
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.paragraphs[0].text = ""
        add_runs(c.paragraphs[0], h, 10.5, base_bold=True)
        shade(c._tc.get_or_add_tcPr(), "E5E7EB")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].paragraphs[0].text = ""
            add_runs(cells[i].paragraphs[0], v, 10.5)
    for r in t.rows:
        for i, c in enumerate(r.cells):
            if widths:
                c.width = Inches(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(2)
    return t


def title_page(doc, meta):
    C = WD_ALIGN_PARAGRAPH.CENTER
    def tp(text, size, bold=False, after=4, italic=False):
        p = para(doc, "", align=C, space_after=after, line=1.15)
        for seg, style in runs(text):
            r = p.add_run(seg)
            set_font(r, FONT, size, bold=bold or style == "bold", italic=italic or style == "italic")
            if style == "fill":
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p
    tp("A REPORT ON", 13, True, 0)
    tp("SUMMER INTERNSHIP", 13, True, 16)
    tp(meta["title"], 19, True, 6)
    tp(meta["subtitle"], 12, italic=True, after=18)
    tp("Submitted in partial fulfilment of the requirements for the award of the degree of", 11.5, after=2)
    tp(meta["programme"], 12.5, True, 14)
    tp("Submitted by", 11.5, after=0)
    tp(meta["student"].upper(), 14, True, 0)
    tp(f"Registration No.: {meta['regno']}", 12, after=12)
    tp("Internship carried out at", 11.5, after=0)
    tp(meta["lab"], 12.5, True, 0)
    tp(meta["org"], 12.5, True, 0)
    tp(meta["period"], 11.5, after=10)
    t = doc.add_table(rows=2, cols=2)
    no_borders(t)
    cells = [("Under the supervision of", f"{meta['supervisor']}\n{meta['supervisor_desig']}\n{meta['org']}"),
             ("Faculty guide", f"{meta['faculty_guide']}\n{meta['faculty_desig']}\n{meta['department']}")]
    for i, (h, body) in enumerate(cells):
        c0 = t.cell(0, i)
        c0.paragraphs[0].alignment = C
        r = c0.paragraphs[0].add_run(h)
        set_font(r, FONT, 11, bold=True)
        c1 = t.cell(1, i)
        p = c1.paragraphs[0]
        p.alignment = C
        p.paragraph_format.line_spacing = 1.15
        for j, line in enumerate(body.split("\n")):
            if j:
                p.add_run().add_break()
            add_runs(p, line, 11)
    para(doc, "", space_after=4)
    placeholder_box(doc, 1.05, "University logo", "Insert the university logo here", width=1.9)
    para(doc, "", space_after=4)
    tp(meta["department"].upper(), 12, True, 0)
    tp(meta["university"].upper(), 12, True, 0)
    tp(meta["campus"], 11.5, after=0)
    tp(meta["month"], 11.5, after=0)


# ------------------------------------------------------------------ main
def render(report, out_path, entries):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin, sec.right_margin = Inches(1.25), Inches(1.0)
    sec.top_margin, sec.bottom_margin = Inches(1.0), Inches(1.0)
    sec.header_distance, sec.footer_distance = Inches(0.5), Inches(0.5)
    setup_styles(doc)

    # front-matter section: roman numbers, nothing on the title page
    sec.different_first_page_header_footer = True
    pgnum(sec, "lowerRoman", 1)
    footer_page(sec)

    tocs = [e for e in entries if e[0] == "toc"]
    figs = [e for e in entries if e[0] == "lof"]
    tabs = [e for e in entries if e[0] == "lot"]

    for b in report.blocks:
        k = b["kind"]
        if k == "title_page":
            title_page(doc, b["meta"])
        elif k == "front":
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.page_break_before = True
            r = p.add_run(b["title"])
            set_font(r, FONT, 18, bold=True)
        elif k in ("toc", "lof", "lot"):
            title = {"toc": "TABLE OF CONTENTS", "lof": "LIST OF FIGURES", "lot": "LIST OF TABLES"}[k]
            if k == "toc":
                p = doc.add_paragraph(style="Front Heading")
                p.paragraph_format.page_break_before = True
                r = p.add_run(title)
                set_font(r, FONT, 18, bold=True)
                contents_list(doc, [e for e in tocs], 'TOC \\o "1-3" \\h \\z \\u',
                              lambda lvl: f"TOC {lvl + 1}")
            else:
                heading1(doc, None, title)
                items = figs if k == "lof" else tabs
                label = "Figure" if k == "lof" else "Table"
                style = "Figure Caption" if k == "lof" else "Table Caption"
                items = [(a, f"{label} {t}", l, pg, kk) for a, t, l, pg, kk in items]
                contents_list(doc, items, f'TOC \\h \\z \\t "{style},1"', lambda lvl: "Table of Figures")
        elif k == "chapter":
            if b["first"]:
                new = doc.add_section(WD_SECTION.NEW_PAGE)
                new.different_first_page_header_footer = False
                pgnum(new, "decimal", 1)
                footer_page(new)
                header_text(new, "Summer Internship Report, IIT Hyderabad", "Diya: Unattended Meditation Kiosk")
                heading1(doc, b["label"], b["title"], page_break=False)
            else:
                heading1(doc, b["label"], b["title"])
        elif k == "section":
            p = doc.add_paragraph(style="Heading 2")
            r = p.add_run(f"{b['num']}   {b['title']}")
            set_font(r, FONT, 14, bold=True)
        elif k == "sub":
            p = doc.add_paragraph(style="Heading 3")
            r = p.add_run(f"{b['num']}  {b['title']}")
            set_font(r, FONT, 12, bold=True)
        elif k == "para":
            para(doc, b["text"], align=WD_ALIGN_PARAGRAPH.LEFT if b["align"] == "left" else None)
        elif k in ("bullets", "numbered"):
            for i, it in enumerate(b["items"], 1):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Inches(0.35)
                pf.first_line_indent = Inches(-0.25)
                pf.space_after = Pt(3)
                pf.tab_stops.add_tab_stop(Inches(0.35))
                r = p.add_run(("•" if k == "bullets" else f"{i}.") + "\t")
                set_font(r, FONT, 12)
                add_runs(p, it)
        elif k == "code":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing = 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.space_after = Pt(10)
            pf.left_indent = Inches(0.05)
            shade(p._p.get_or_add_pPr(), "F3F4F6")
            for j, line in enumerate(b["text"].split("\n")):
                r = p.add_run(line)
                set_font(r, MONO, 9)
                if j < len(b["text"].split("\n")) - 1:
                    r.add_break()
        elif k == "figure":
            picture_par(doc, b["path"], b["width"])
            caption_par(doc, "Figure", b["num"], b["caption"], "Figure Caption")
        elif k == "figrow":
            n = len(b["items"])
            t = doc.add_table(rows=2, cols=n)
            no_borders(t)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (path, lab) in enumerate(b["items"]):
                c = t.cell(0, i)
                c.width = Inches(TEXT_W / n)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                with PILImage.open(path) as im:
                    w, h = im.size
                width = min(b["height"] * w / h, TEXT_W / n - 0.1)
                p.add_run().add_picture(path, width=Inches(width))
                lp = t.cell(1, i).paragraphs[0]
                lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                lp.paragraph_format.keep_with_next = True
                add_runs(lp, lab, 10)
            caption_par(doc, "Figure", b["num"], b["caption"], "Figure Caption")
        elif k == "placeholder":
            title = "PLACE PHOTOGRAPH HERE" if b["num"] else "PLACE SCANNED CERTIFICATE HERE"
            placeholder_box(doc, b["height"], title, b["hint"])
            if b["num"]:
                caption_par(doc, "Figure", b["num"], b["caption"], "Figure Caption")
        elif k == "table":
            if b["num"]:
                caption_par(doc, "Table", b["num"], b["caption"], "Table Caption")
            data_table(doc, b["header"], b["rows"], b["widths"])
            para(doc, "", space_after=4)
        elif k == "signature":
            for ln in ["(Signature)"] + b["lines"]:
                para(doc, ln, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
        elif k == "refs":
            for i, it in enumerate(b["items"], 1):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Inches(0.4)
                pf.first_line_indent = Inches(-0.4)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.line_spacing = 1.3
                pf.tab_stops.add_tab_stop(Inches(0.4))
                r = p.add_run(f"[{i}]\t")
                set_font(r, FONT, 11.5)
                add_runs(p, it, 11.5)
        elif k == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif k == "spacer":
            para(doc, "", space_after=b["h"] * 72)

    # ask Word to refresh the contents fields when the file is opened
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    insert_ordered(settings, uf, SETTINGS_AFTER_UPDATE)
    core = doc.core_properties
    core.title = "Diya: Design and Development of an Unattended Meditation Kiosk"
    core.author = "Ayush Gupta"
    core.subject = "Summer Internship Report, IIT Hyderabad"
    doc.save(out_path)
    print(f"DOCX: {out_path}")
